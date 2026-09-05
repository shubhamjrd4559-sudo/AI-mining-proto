"""
apps.maintainer — Views (Phase 4)

All endpoints require authentication.
All dataset access is owner-scoped through source_document.uploaded_by.

Endpoints:
  GET  /api/maintainer/datasets/
  GET  /api/maintainer/datasets/<id>/overview/
  GET  /api/maintainer/datasets/<id>/records/
  GET  /api/maintainer/datasets/<id>/suggestions/
  POST /api/maintainer/datasets/<id>/generate-suggestions/
  POST /api/maintainer/suggestions/<id>/approve/
  POST /api/maintainer/suggestions/<id>/reject/
  POST /api/maintainer/suggestions/batch-review/
  POST /api/maintainer/datasets/<id>/apply/
  GET  /api/maintainer/datasets/<id>/export/xlsx/
  GET  /api/maintainer/datasets/<id>/export/csv/
"""

import csv
import io
import logging
from django.db import transaction, DatabaseError
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.audit.models import AuditEvent, AuditEventType
from apps.datasets.models import StructuredDataset, StructuredRecord
from apps.pipeline.models import ValidationResult

from .models import MaintainerSuggestion, SuggestionStatus
from .serializers import (
    DatasetOverviewSerializer,
    MaintainerSuggestionSerializer,
    StructuredRecordSerializer,
)
from .suggester import generate_suggestions_for_dataset

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────

def _get_client_ip(request) -> str | None:
    xff = request.META.get('HTTP_X_FORWARDED_FOR')
    if xff:
        return xff.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR')


def _get_owned_dataset(dataset_id, request) -> tuple:
    """
    Returns (dataset, error_response).
    Enforces that the requesting user owns the dataset (via source_document.uploaded_by).
    Returns (None, Response) on failure.
    """
    try:
        dataset = StructuredDataset.objects.select_related('source_document').get(pk=dataset_id)
    except StructuredDataset.DoesNotExist:
        return None, Response({'error': 'Dataset not found.'}, status=status.HTTP_404_NOT_FOUND)

    doc = dataset.source_document
    if not doc or not doc.uploaded_by or doc.uploaded_by != request.user:
        return None, Response(
            {'error': 'You do not have permission to access this dataset.'},
            status=status.HTTP_403_FORBIDDEN,
        )
    return dataset, None


def _get_owned_suggestion(suggestion_id, request) -> tuple:
    """Returns (suggestion, error_response) with ownership check."""
    try:
        suggestion = MaintainerSuggestion.objects.select_related('dataset__source_document').get(pk=suggestion_id)
    except MaintainerSuggestion.DoesNotExist:
        return None, Response({'error': 'Suggestion not found.'}, status=status.HTTP_404_NOT_FOUND)

    doc = suggestion.dataset.source_document if suggestion.dataset else None
    if not doc or not doc.uploaded_by or doc.uploaded_by != request.user:
        return None, Response(
            {'error': 'You do not have permission to access this suggestion.'},
            status=status.HTTP_403_FORBIDDEN,
        )
    return suggestion, None


def _log_audit(request, event_type: str, resource_id: str, description: str,
               metadata: dict = None) -> None:
    """Create an AuditEvent — call inside a transaction."""
    actor = request.user.username if request.user and request.user.is_authenticated else 'anonymous'
    AuditEvent.objects.create(
        event_type=event_type,
        actor=actor,
        actor_ip=_get_client_ip(request),
        resource_type='maintainer_suggestion',
        resource_id=str(resource_id),
        description=description,
        metadata_json=metadata or {},
    )


# ─────────────────────────────────────────────
# Dataset list
# ─────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dataset_list(request):
    """GET /api/maintainer/datasets/ — list datasets owned by requesting user."""
    qs = StructuredDataset.objects.filter(
        source_document__uploaded_by=request.user
    ).select_related('source_document').order_by('-created_at')

    paginator = PageNumberPagination()
    page = paginator.paginate_queryset(qs, request)
    if page is not None:
        serializer = DatasetOverviewSerializer(page, many=True)
        return paginator.get_paginated_response(serializer.data)

    serializer = DatasetOverviewSerializer(qs, many=True)
    return Response({'count': qs.count(), 'results': serializer.data})


# ─────────────────────────────────────────────
# Dataset overview
# ─────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dataset_overview(request, dataset_id):
    """GET /api/maintainer/datasets/<id>/overview/"""
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    serializer = DatasetOverviewSerializer(dataset)
    data = serializer.data

    # Append suggestion summary
    suggestions_qs = MaintainerSuggestion.objects.filter(dataset=dataset)
    data['suggestion_summary'] = {
        'pending': suggestions_qs.filter(status=SuggestionStatus.PENDING).count(),
        'approved': suggestions_qs.filter(status=SuggestionStatus.APPROVED).count(),
        'rejected': suggestions_qs.filter(status=SuggestionStatus.REJECTED).count(),
        'applied': suggestions_qs.filter(status=SuggestionStatus.APPLIED).count(),
        'total': suggestions_qs.count(),
    }

    return Response(data)


# ─────────────────────────────────────────────
# Records (preview)
# ─────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dataset_records(request, dataset_id):
    """GET /api/maintainer/datasets/<id>/records/"""
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    qs = dataset.records.all().order_by('row_index').prefetch_related('provenance', 'suggestions')

    paginator = PageNumberPagination()
    paginator.page_size = request.query_params.get('page_size', 50)
    page = paginator.paginate_queryset(qs, request)
    if page is not None:
        serializer = StructuredRecordSerializer(page, many=True)
        return paginator.get_paginated_response(serializer.data)

    serializer = StructuredRecordSerializer(qs, many=True)
    return Response({'count': qs.count(), 'results': serializer.data})


# ─────────────────────────────────────────────
# Suggestions list
# ─────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dataset_suggestions(request, dataset_id):
    """GET /api/maintainer/datasets/<id>/suggestions/?status=&issue_type="""
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    qs = MaintainerSuggestion.objects.filter(dataset=dataset).select_related(
        'record', 'created_by', 'reviewed_by'
    ).order_by('-created_at')

    status_filter = request.query_params.get('status')
    if status_filter:
        qs = qs.filter(status=status_filter)

    issue_filter = request.query_params.get('issue_type')
    if issue_filter:
        qs = qs.filter(issue_type=issue_filter)

    paginator = PageNumberPagination()
    page = paginator.paginate_queryset(qs, request)
    if page is not None:
        serializer = MaintainerSuggestionSerializer(page, many=True)
        return paginator.get_paginated_response(serializer.data)

    serializer = MaintainerSuggestionSerializer(qs, many=True)
    return Response({'count': qs.count(), 'results': serializer.data})


# ─────────────────────────────────────────────
# Generate suggestions
# ─────────────────────────────────────────────

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def generate_suggestions(request, dataset_id):
    """POST /api/maintainer/datasets/<id>/generate-suggestions/"""
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    if dataset.record_count == 0 and not dataset.records.exists():
        return Response(
            {'error': 'Dataset has no records. Run document processing first.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    import os
    ai_enabled = bool(
        os.environ.get('OPENAI_API_KEY', '').strip() or
        os.environ.get('GEMINI_API_KEY', '').strip()
    )

    try:
        summary = generate_suggestions_for_dataset(dataset, user=request.user, ai_enabled=ai_enabled)
    except Exception as exc:
        logger.error('Suggestion generation failed for dataset %s: %s', dataset_id, exc, exc_info=True)
        return Response(
            {'error': f'Suggestion generation failed: {exc}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    return Response({
        'status': 'ok',
        'dataset_id': dataset_id,
        'ai_enabled': ai_enabled,
        **summary,
    }, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────
# Approve single suggestion
# ─────────────────────────────────────────────

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def approve_suggestion(request, suggestion_id):
    """POST /api/maintainer/suggestions/<id>/approve/
    APPROVE DOES NOT APPLY — only changes status to APPROVED.
    """
    suggestion, err = _get_owned_suggestion(suggestion_id, request)
    if err:
        return err

    if suggestion.status != SuggestionStatus.PENDING:
        return Response(
            {'error': f'Cannot approve suggestion in "{suggestion.status}" status. Only PENDING suggestions can be approved.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    now = timezone.now()
    with transaction.atomic():
        suggestion.status = SuggestionStatus.APPROVED
        suggestion.reviewed_by = request.user
        suggestion.reviewed_at = now
        suggestion.save(update_fields=['status', 'reviewed_by', 'reviewed_at'])

        AuditEvent.objects.create(
            event_type=AuditEventType.DOCUMENT_REVIEWED,
            actor=request.user.username,
            actor_ip=_get_client_ip(request),
            resource_type='maintainer_suggestion',
            resource_id=str(suggestion.pk),
            description=(
                f'Approved suggestion {suggestion.pk} for field "{suggestion.field_name}" '
                f'in dataset "{suggestion.dataset.name}"'
            ),
            metadata_json={
                'dataset_id': suggestion.dataset.pk,
                'suggestion_id': suggestion.pk,
                'action': 'approve',
                'field_name': suggestion.field_name,
                'original_value': suggestion.original_value,
                'suggested_value': suggestion.suggested_value,
            },
        )

    serializer = MaintainerSuggestionSerializer(suggestion)
    return Response(serializer.data, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────
# Reject single suggestion
# ─────────────────────────────────────────────

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def reject_suggestion(request, suggestion_id):
    """POST /api/maintainer/suggestions/<id>/reject/"""
    suggestion, err = _get_owned_suggestion(suggestion_id, request)
    if err:
        return err

    if suggestion.status not in (SuggestionStatus.PENDING, SuggestionStatus.APPROVED):
        return Response(
            {'error': f'Cannot reject suggestion in "{suggestion.status}" status.'},
            status=status.HTTP_400_BAD_REQUEST,
        )

    now = timezone.now()
    with transaction.atomic():
        suggestion.status = SuggestionStatus.REJECTED
        suggestion.reviewed_by = request.user
        suggestion.reviewed_at = now
        suggestion.save(update_fields=['status', 'reviewed_by', 'reviewed_at'])

        AuditEvent.objects.create(
            event_type=AuditEventType.DOCUMENT_REVIEWED,
            actor=request.user.username,
            actor_ip=_get_client_ip(request),
            resource_type='maintainer_suggestion',
            resource_id=str(suggestion.pk),
            description=(
                f'Rejected suggestion {suggestion.pk} for field "{suggestion.field_name}" '
                f'in dataset "{suggestion.dataset.name}"'
            ),
            metadata_json={
                'dataset_id': suggestion.dataset.pk,
                'suggestion_id': suggestion.pk,
                'action': 'reject',
                'field_name': suggestion.field_name,
                'original_value': suggestion.original_value,
                'suggested_value': suggestion.suggested_value,
            },
        )

    serializer = MaintainerSuggestionSerializer(suggestion)
    return Response(serializer.data, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────
# Batch approve / reject
# ─────────────────────────────────────────────

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def batch_review(request):
    """POST /api/maintainer/suggestions/batch-review/
    Body: {"ids": [1,2,3], "action": "approve"|"reject"}
    """
    ids = request.data.get('ids', [])
    action = request.data.get('action', '')

    if not ids:
        return Response({'error': 'No suggestion ids provided.'}, status=status.HTTP_400_BAD_REQUEST)
    if action not in ('approve', 'reject'):
        return Response({'error': 'action must be "approve" or "reject".'}, status=status.HTTP_400_BAD_REQUEST)
    if len(ids) > 500:
        return Response({'error': 'Maximum 500 suggestions per batch.'}, status=status.HTTP_400_BAD_REQUEST)

    # Verify all suggestions belong to this user
    suggestions = list(MaintainerSuggestion.objects.filter(pk__in=ids).select_related(
        'dataset__source_document'
    ))

    # Owner check
    for s in suggestions:
        doc = s.dataset.source_document if s.dataset else None
        if not doc or not doc.uploaded_by or doc.uploaded_by != request.user:
            return Response(
                {'error': f'Suggestion {s.pk} does not belong to you.'},
                status=status.HTTP_403_FORBIDDEN,
            )

    new_status = SuggestionStatus.APPROVED if action == 'approve' else SuggestionStatus.REJECTED

    if action == 'approve':
        valid_statuses = [SuggestionStatus.PENDING]
    else:
        valid_statuses = [SuggestionStatus.PENDING, SuggestionStatus.APPROVED]

    now = timezone.now()
    with transaction.atomic():
        updated = MaintainerSuggestion.objects.filter(
            pk__in=ids, status__in=valid_statuses
        ).update(
            status=new_status,
            reviewed_by=request.user,
            reviewed_at=now,
        )

        for s in suggestions:
            if s.status in valid_statuses:
                AuditEvent.objects.create(
                    event_type=AuditEventType.DOCUMENT_REVIEWED,
                    actor=request.user.username,
                    actor_ip=_get_client_ip(request),
                    resource_type='maintainer_suggestion',
                    resource_id=str(s.pk),
                    description=(
                        f'Batch {action} suggestion {s.pk} for field "{s.field_name}" '
                        f'in dataset "{s.dataset.name}"'
                    ),
                    metadata_json={
                        'dataset_id': s.dataset.pk,
                        'suggestion_id': s.pk,
                        'action': action,
                        'field_name': s.field_name,
                        'original_value': s.original_value,
                        'suggested_value': s.suggested_value,
                    },
                )

    return Response({
        'action': action,
        'updated': updated,
        'requested': len(ids),
        'skipped': len(ids) - updated,
    }, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────
# Apply approved suggestions
# ─────────────────────────────────────────────

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def apply_approved(request, dataset_id):
    """
    POST /api/maintainer/datasets/<id>/apply/
    Applies all APPROVED suggestions to their StructuredRecord.data_json.
    Uses transaction.atomic(). Prevents duplicate application.
    Enforces conflict check: record current field value must match suggestion.original_value.
    Original uploaded file is NEVER touched.
    """
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    approved = MaintainerSuggestion.objects.filter(
        dataset=dataset,
        status=SuggestionStatus.APPROVED,
    ).select_related('record').order_by('id')

    if not approved.exists():
        return Response(
            {'message': 'No approved suggestions to apply.', 'applied': 0, 'failed': 0},
            status=status.HTTP_200_OK,
        )

    applied_count = 0
    failed_count = 0
    failed_ids = []
    now = timezone.now()

    for suggestion in approved:
        if not suggestion.record:
            suggestion.status = SuggestionStatus.FAILED
            suggestion.error_message = 'No associated record — cannot apply.'
            suggestion.save(update_fields=['status', 'error_message'])
            failed_count += 1
            failed_ids.append(suggestion.pk)
            continue

        try:
            with transaction.atomic():
                # Re-fetch record with select_for_update to prevent concurrent modification
                record = StructuredRecord.objects.select_for_update().get(pk=suggestion.record_id)

                # Re-fetch suggestion inside transaction to prevent race condition on duplicate apply
                locked_suggestion = MaintainerSuggestion.objects.select_for_update().get(pk=suggestion.pk)
                if locked_suggestion.status != SuggestionStatus.APPROVED:
                    # Already applied/failed by a concurrent request
                    continue

                row = dict(record.data_json or {})
                field = locked_suggestion.field_name

                # Conflict Check: current value must match suggestion.original_value
                current_val = str(row.get(field, '')) if row.get(field) is not None else ''
                expected_val = str(locked_suggestion.original_value) if locked_suggestion.original_value is not None else ''

                if current_val != expected_val:
                    locked_suggestion.status = SuggestionStatus.FAILED
                    locked_suggestion.error_message = (
                        f'Conflict detected: current value "{current_val}" '
                        f'does not match expected original value "{expected_val}".'
                    )
                    locked_suggestion.save(update_fields=['status', 'error_message'])
                    failed_count += 1
                    failed_ids.append(locked_suggestion.pk)
                    continue

                # Store original before modification in non-destructive key
                old_val = row.get(field)
                orig_key = f'{field}_original_before_apply'
                if orig_key not in row:
                    row[orig_key] = old_val

                # Apply correction
                row[field] = locked_suggestion.suggested_value

                record.data_json = row
                record.save(update_fields=['data_json'])

                # Update suggestion
                locked_suggestion.status = SuggestionStatus.APPLIED
                locked_suggestion.applied_value = locked_suggestion.suggested_value
                locked_suggestion.applied_at = now
                locked_suggestion.save(update_fields=['status', 'applied_value', 'applied_at'])

                # AuditEvent
                AuditEvent.objects.create(
                    event_type=AuditEventType.DATASET_UPDATED,
                    actor=request.user.username,
                    actor_ip=_get_client_ip(request),
                    resource_type='maintainer_suggestion',
                    resource_id=str(locked_suggestion.pk),
                    description=(
                        f'Applied correction to dataset "{dataset.name}" '
                        f'record {record.row_index}, field "{field}": '
                        f'"{old_val}" → "{locked_suggestion.suggested_value}"'
                    ),
                    metadata_json={
                        'dataset_id': dataset.pk,
                        'record_id': str(record.pk),
                        'record_row_index': record.row_index,
                        'field_name': field,
                        'old_value': str(old_val),
                        'new_value': locked_suggestion.suggested_value,
                        'suggestion_id': locked_suggestion.pk,
                        'suggestion_source': locked_suggestion.suggestion_source,
                        'issue_type': locked_suggestion.issue_type,
                        'confidence': locked_suggestion.confidence,
                        'reason': locked_suggestion.reason,
                    },
                )

                applied_count += 1

        except StructuredRecord.DoesNotExist:
            suggestion.status = SuggestionStatus.FAILED
            suggestion.error_message = 'Record no longer exists.'
            suggestion.save(update_fields=['status', 'error_message'])
            failed_count += 1
            failed_ids.append(suggestion.pk)
        except (DatabaseError, Exception) as exc:
            logger.error('Apply failed for suggestion %s: %s', suggestion.pk, exc, exc_info=True)
            try:
                MaintainerSuggestion.objects.filter(pk=suggestion.pk).update(
                    status=SuggestionStatus.FAILED,
                    error_message=str(exc)[:500],
                )
            except Exception:
                pass
            failed_count += 1
            failed_ids.append(suggestion.pk)

    return Response({
        'status': 'ok',
        'applied': applied_count,
        'failed': failed_count,
        'failed_ids': failed_ids,
    }, status=status.HTTP_200_OK)


# ─────────────────────────────────────────────
# XLSX Export
# ─────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_xlsx(request, dataset_id):
    """
    GET /api/maintainer/datasets/<id>/export/xlsx/
    Exports currently applied (corrected) record values as XLSX.
    Original source file is NEVER modified.
    """
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return Response({'error': 'openpyxl not installed.'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Maintained Data'

    records = list(dataset.records.order_by('row_index'))
    if not records:
        ws.cell(row=1, column=1, value='No records found.')
        buf = io.BytesIO()
        wb.save(buf)
        content = buf.getvalue()
        resp = HttpResponse(
            content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        safe_name = _safe_export_name(dataset)
        resp['Content-Disposition'] = f'attachment; filename="{safe_name}.xlsx"'
        return resp

    # Build headers from schema + all keys in records
    schema_cols = list((dataset.schema_json or {}).get('columns', []))
    all_keys: list = []
    seen = set()
    for col in schema_cols:
        if col not in seen:
            all_keys.append(col)
            seen.add(col)
    for record in records:
        for k in (record.data_json or {}):
            if k not in seen and not k.endswith('_original_before_apply'):
                all_keys.append(k)
                seen.add(k)

    # Header row
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    for col_idx, header in enumerate(all_keys, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font

    # Data rows
    applied_fields = _get_applied_fields(dataset)
    correction_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')

    for row_idx, record in enumerate(records, start=2):
        row_data = record.data_json or {}
        for col_idx, header in enumerate(all_keys, start=1):
            val = row_data.get(header, '')
            cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else '')
            # Highlight corrected cells
            if (record.pk, header) in applied_fields:
                cell.fill = correction_fill

    # Metadata sheet
    ws_meta = wb.create_sheet(title='Maintainer Info')
    ws_meta.cell(row=1, column=1, value='Dataset')
    ws_meta.cell(row=1, column=2, value=dataset.name)
    ws_meta.cell(row=2, column=1, value='Source Document')
    ws_meta.cell(row=2, column=2, value=dataset.source_document.original_filename if dataset.source_document else '')
    ws_meta.cell(row=3, column=1, value='Records')
    ws_meta.cell(row=3, column=2, value=len(records))
    ws_meta.cell(row=4, column=1, value='Exported At')
    ws_meta.cell(row=4, column=2, value=timezone.now().isoformat())
    ws_meta.cell(row=5, column=1, value='Note')
    ws_meta.cell(row=5, column=2, value='This is a DERIVED file. Original source is immutable.')

    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    safe_name = _safe_export_name(dataset)
    resp = HttpResponse(
        content,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = f'attachment; filename="{safe_name}_maintained.xlsx"'
    resp['Content-Length'] = len(content)
    return resp


# ─────────────────────────────────────────────
# CSV Export
# ─────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_csv(request, dataset_id):
    """
    GET /api/maintainer/datasets/<id>/export/csv/
    Exports corrected data as UTF-8 CSV.
    Original source is NEVER modified.
    """
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    records = list(dataset.records.order_by('row_index'))

    # Build headers
    schema_cols = list((dataset.schema_json or {}).get('columns', []))
    all_keys: list = []
    seen = set()
    for col in schema_cols:
        if col not in seen:
            all_keys.append(col)
            seen.add(col)
    for record in records:
        for k in (record.data_json or {}):
            if k not in seen and not k.endswith('_original_before_apply'):
                all_keys.append(k)
                seen.add(k)

    output = io.StringIO()
    writer = csv.writer(output, lineterminator='\n')
    writer.writerow(all_keys)

    for record in records:
        row_data = record.data_json or {}
        writer.writerow([str(row_data.get(h, '')) if row_data.get(h) is not None else '' for h in all_keys])

    safe_name = _safe_export_name(dataset)
    resp = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8')
    resp['Content-Disposition'] = f'attachment; filename="{safe_name}_maintained.csv"'
    return resp


# ─────────────────────────────────────────────
# Validation issues (for frontend display)
# ─────────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dataset_validation_issues(request, dataset_id):
    """GET /api/maintainer/datasets/<id>/validation-issues/"""
    dataset, err = _get_owned_dataset(dataset_id, request)
    if err:
        return err

    if not dataset.source_document:
        return Response({'count': 0, 'results': []})

    qs = ValidationResult.objects.filter(
        document=dataset.source_document
    ).order_by('-severity', 'issue_type')

    status_filter = request.query_params.get('status')
    if status_filter:
        qs = qs.filter(status=status_filter)

    issue_filter = request.query_params.get('issue_type')
    if issue_filter:
        qs = qs.filter(issue_type=issue_filter)

    paginator = PageNumberPagination()
    page = paginator.paginate_queryset(qs, request)

    def _ser(v):
        return {
            'id': v.pk,
            'field_name': v.field_name,
            'issue_type': v.issue_type,
            'severity': v.severity,
            'original_value': v.original_value,
            'suggested_value': v.suggested_value,
            'explanation': v.explanation,
            'confidence': v.confidence,
            'status': v.status,
            'created_at': v.created_at.isoformat(),
        }

    if page is not None:
        return paginator.get_paginated_response([_ser(v) for v in page])

    return Response({'count': qs.count(), 'results': [_ser(v) for v in qs]})


# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────

def _get_applied_fields(dataset) -> set:
    """Return set of (record_id, field_name) tuples for applied suggestions."""
    applied = MaintainerSuggestion.objects.filter(
        dataset=dataset, status=SuggestionStatus.APPLIED
    ).values_list('record_id', 'field_name')
    return set(applied)


def _safe_export_name(dataset) -> str:
    """Generate a safe filename for export."""
    import re
    raw = dataset.source_document.original_filename if dataset.source_document else dataset.name
    # Strip extension
    name = raw.rsplit('.', 1)[0] if '.' in raw else raw
    safe = re.sub(r'[^a-zA-Z0-9_\-]', '_', name)[:60]
    return safe or 'dataset'
