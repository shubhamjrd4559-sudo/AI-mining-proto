import io
import struct
import threading
from django.test import TestCase
from django.contrib.auth import get_user_model
from rest_framework.test import APIClient
from rest_framework import status

from apps.documents.models import Document, DocumentStatus, ProcessingJob
from apps.datasets.models import StructuredDataset, StructuredRecord
from apps.storage.service import get_storage_service
from .models import ExtractionResult, ExtractionProvenance, ValidationResult
from .extractor.pdf_extractor import extract_pdf
from .extractor.docx_extractor import extract_docx
from .extractor.xlsx_extractor import extract_xlsx
from .extractor.csv_extractor import extract_csv
from .extractor.txt_extractor import extract_txt
from .extractor.image_extractor import extract_image
from .schema.detector import map_columns, detect_field
from .schema.normalizer import normalize_value, normalize_row
from .validation.engine import validate_dataset
from .extractor.base import ExtractedTable
from .orchestrator import _run_pipeline

User = get_user_model()

MINIMAL_PDF = (
    b'%PDF-1.4\n'
    b'1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n'
    b'2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n'
    b'3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]\n'
    b'  /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n'
    b'4 0 obj\n<< /Length 44 >>\nstream\nBT /F1 12 Tf 100 700 Td (Mine: Jharia) Tj ET\nendstream\nendobj\n'
    b'5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n'
    b'xref\n0 6\n0000000000 65535 f \ntrailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n9\n%%EOF'
)

def _make_xlsx_bytes():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Production'
    ws.append(['Mine', 'Production (MT)', 'FY', 'Subsidiary'])
    ws.append(['Jharia', '4.5 MT', '2024-25', 'BCCL'])
    ws.append(['Kathara', '2.1 MT', '2024-25', 'CCL'])
    
    ws2 = wb.create_sheet(title='Targets')
    ws2.append(['Mine', 'Target (MT)'])
    ws2.append(['Jharia', '5.0'])
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def _make_docx_bytes():
    from docx import Document
    doc = Document()
    doc.add_heading('Mining Report', 0)
    doc.add_paragraph('Production data for FY 2024-25.')
    tbl = doc.add_table(rows=2, cols=3)
    tbl.rows[0].cells[0].text = 'Mine'
    tbl.rows[0].cells[1].text = 'Production'
    tbl.rows[0].cells[2].text = 'FY'
    tbl.rows[1].cells[0].text = 'Jharia'
    tbl.rows[1].cells[1].text = '4.5 MT'
    tbl.rows[1].cells[2].text = '2024-25'
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def _make_png_bytes():
    # 1x1 black pixel PNG
    return b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\x0bIDAT\x08\x99c\xf8\x0f\x04\x00\t\xfb\x03\xfd\xe3U\xf2\x9c\x00\x00\x00\x00IEND\xaeB`\x82'


class ExtractorTests(TestCase):
    def test_pdf_extraction(self):
        doc = extract_pdf(MINIMAL_PDF)
        self.assertEqual(doc.extractor_type, 'pdf_text')
        self.assertFalse(doc.ocr_used)
        
    def test_pdf_ocr_fallback(self):
        # Image-only PDF will trigger OCR if pytesseract is available.
        # Since we use a minimal pdf without text layer but minimal chars, it may trigger ocr.
        pass

    def test_docx_paragraphs(self):
        doc = extract_docx(_make_docx_bytes())
        self.assertEqual(doc.extractor_type, 'docx')
        self.assertIn('Mining Report', doc.raw_text)

    def test_docx_tables(self):
        doc = extract_docx(_make_docx_bytes())
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].headers, ['Mine', 'Production', 'FY'])
        self.assertEqual(doc.tables[0].rows[0], ['Jharia', '4.5 MT', '2024-25'])

    def test_xlsx_headers_rows(self):
        doc = extract_xlsx(_make_xlsx_bytes())
        self.assertEqual(doc.extractor_type, 'xlsx')
        self.assertEqual(len(doc.tables), 2)
        self.assertEqual(doc.tables[0].headers, ['Mine', 'Production (MT)', 'FY', 'Subsidiary'])
        self.assertEqual(doc.tables[0].rows[0], ['Jharia', '4.5 MT', '2024-25', 'BCCL'])

    def test_xlsx_multisheet(self):
        doc = extract_xlsx(_make_xlsx_bytes())
        self.assertEqual(doc.tables[1].sheet_name, 'Targets')

    def test_csv_comma(self):
        csv_data = b'Mine,Production\nJharia,4.5\n'
        doc = extract_csv(csv_data)
        self.assertEqual(doc.tables[0].headers, ['Mine', 'Production'])
        self.assertEqual(doc.tables[0].rows[0], ['Jharia', '4.5'])
        
    def test_csv_semicolon(self):
        csv_data = b'Mine;Production\nJharia;4.5\n'
        doc = extract_csv(csv_data)
        self.assertEqual(doc.tables[0].headers, ['Mine', 'Production'])

    def test_txt_utf8(self):
        doc = extract_txt(b'Hello World')
        self.assertEqual(doc.raw_text, 'Hello World')

    def test_txt_latin1(self):
        doc = extract_txt(b'Hello \xff')
        self.assertIn('Hello', doc.raw_text)

    def test_image_ocr(self):
        doc = extract_image(_make_png_bytes())
        self.assertEqual(doc.extractor_type, 'image_ocr')
        # If tesseract unavailable, returns ocr_unavailable in error/status_note


class SchemaTests(TestCase):
    def test_detector_production(self):
        self.assertEqual(detect_field('Production (MT)'), 'production')

    def test_detector_fy(self):
        self.assertEqual(detect_field('Fin Year'), 'financial_year')

    def test_detector_unknown(self):
        self.assertIsNone(detect_field('Random Col'))

    def test_normalizer_fy(self):
        # FY 2024-25 has a separator and should normalize correctly
        res = normalize_value('FY 2024-25', 'financial_year')
        self.assertEqual(res['normalized_value'], '2024-25')

    def test_normalizer_numeric(self):
        res = normalize_value('4.5 MT', 'production')
        self.assertEqual(res['normalized_value'], 4.5)
        self.assertEqual(res['normalized_unit'], 'MT')


class ValidationTests(TestCase):
    def test_missing_important(self):
        t1 = ExtractedTable(source_ref='t1', headers=['Random'], rows=[['1']])
        findings = validate_dataset([t1], map_columns(['Random']))
        issues = [f.issue_type for f in findings]
        self.assertIn('missing_required', issues)

    def test_duplicate_column(self):
        t1 = ExtractedTable(source_ref='t1', headers=['Col', 'Col'], rows=[['1', '2']])
        findings = validate_dataset([t1], map_columns(['Col']))
        self.assertTrue(any(f.issue_type == 'duplicate' for f in findings))

    def test_invalid_numeric(self):
        t1 = ExtractedTable(source_ref='t1', headers=['Production'], rows=[['N/A']])
        findings = validate_dataset([t1], map_columns(['Production']))
        self.assertTrue(any(f.issue_type == 'invalid_numeric' for f in findings))

    def test_suspicious_negative(self):
        t1 = ExtractedTable(source_ref='t1', headers=['Production'], rows=[['-5.0']])
        findings = validate_dataset([t1], map_columns(['Production']))
        self.assertTrue(any(f.issue_type == 'suspicious_value' for f in findings))

    def test_coordinate_out_of_range(self):
        t1 = ExtractedTable(source_ref='t1', headers=['Latitude'], rows=[['95.0']])
        findings = validate_dataset([t1], map_columns(['Latitude']))
        self.assertTrue(any(f.issue_type == 'coordinate_error' for f in findings))


class OrchestratorTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='tester', password='pw')
        self.doc = Document.objects.create(
            title='Test',
            original_filename='test.csv',
            stored_filename='test.csv',
            storage_key='documents/test.csv',
            file_extension='.csv',
            status=DocumentStatus.UPLOADED,
            uploaded_by=self.user
        )
        storage = get_storage_service()
        storage.save(self.doc.storage_key, b'Mine,Production\nJharia,4.5\n')

    def test_pipeline_status_transitions(self):
        _run_pipeline(self.doc.pk)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.status, DocumentStatus.INDEXED)
        self.assertTrue(ExtractionResult.objects.filter(document=self.doc).exists())
        self.assertTrue(StructuredDataset.objects.filter(source_document=self.doc).exists())

    def test_pipeline_failed_extraction(self):
        self.doc.file_extension = '.unknown'
        self.doc.save()
        _run_pipeline(self.doc.pk)
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.status, DocumentStatus.FAILED)


class ConcurrencyAndRetryTests(TestCase):
    def test_transient_lock_detection(self):
        from django.db import OperationalError
        from .orchestrator import is_transient_db_error

        self.assertTrue(is_transient_db_error(OperationalError('database table is locked: documents_document')))
        self.assertTrue(is_transient_db_error(OperationalError('database is locked')))
        self.assertTrue(is_transient_db_error(OperationalError('sqlite3.OperationalError: database is busy')))
        self.assertFalse(is_transient_db_error(ValueError('random error')))
        self.assertFalse(is_transient_db_error(OperationalError('syntax error near SELECT')))

    def test_db_retry_recovers_after_transient_lock(self):
        from django.db import OperationalError
        from .orchestrator import db_retry

        attempts = 0

        @db_retry(max_retries=3, initial_delay=0.01, max_delay=0.05)
        def flaky_operation():
            nonlocal attempts
            attempts += 1
            if attempts < 3:
                raise OperationalError('database is locked')
            return 'success'

        result = flaky_operation()
        self.assertEqual(result, 'success')
        self.assertEqual(attempts, 3)

    def test_db_retry_exhausted_raises_exception(self):
        from django.db import OperationalError
        from .orchestrator import db_retry

        @db_retry(max_retries=2, initial_delay=0.01, max_delay=0.02)
        def always_locked():
            raise OperationalError('database is locked')

        with self.assertRaises(OperationalError):
            always_locked()

    def test_db_retry_does_not_retry_non_transient_error(self):
        from .orchestrator import db_retry

        attempts = 0

        @db_retry(max_retries=3, initial_delay=0.01, max_delay=0.05)
        def bad_code():
            nonlocal attempts
            attempts += 1
            raise ValueError('bad parameter')

        with self.assertRaises(ValueError):
            bad_code()
        self.assertEqual(attempts, 1)


class APIExtractionTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='tester', password='pw')
        self.other = User.objects.create_user(username='other', password='pw')
        self.client = APIClient()
        
        self.doc = Document.objects.create(
            title='Test', original_filename='test.csv',
            status=DocumentStatus.UPLOADED, uploaded_by=self.user
        )
        self.er = ExtractionResult.objects.create(
            document=self.doc, extractor_type='csv', status='completed'
        )

    def test_requires_auth(self):
        resp = self.client.get(f'/api/documents/{self.doc.pk}/extraction/')
        self.assertEqual(resp.status_code, 401)

    def test_owner_access(self):
        self.client.force_authenticate(user=self.user)
        resp = self.client.get(f'/api/documents/{self.doc.pk}/extraction/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['extractor_type'], 'csv')

    def test_other_user_denied(self):
        self.client.force_authenticate(user=self.other)
        resp = self.client.get(f'/api/documents/{self.doc.pk}/extraction/')
        self.assertEqual(resp.status_code, 404)
