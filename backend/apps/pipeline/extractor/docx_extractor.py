"""
DOCX extractor using python-docx.
"""
import io
import logging
from .base import ExtractedDocument, ExtractedTable

logger = logging.getLogger(__name__)

HEADING_STYLES = {'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4',
                   'Title', 'Subtitle', 'heading 1', 'heading 2', 'heading 3'}


def extract_docx(file_bytes: bytes) -> ExtractedDocument:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return ExtractedDocument(extractor_type='docx', error='python-docx not installed')

    try:
        doc = DocxDocument(io.BytesIO(file_bytes))
    except Exception as exc:
        return ExtractedDocument(extractor_type='docx', error=f'Cannot open DOCX: {exc}')

    paragraphs = []
    current_heading = ''
    tables = []

    try:
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ''
            if style_name in HEADING_STYLES:
                current_heading = text
            paragraphs.append(text)

        for t_idx, tbl in enumerate(doc.tables):
            headers = []
            rows = []
            for r_idx, row in enumerate(tbl.rows):
                cells = [cell.text.strip() for cell in row.cells]
                if r_idx == 0:
                    headers = cells
                else:
                    rows.append(cells)
            tables.append(ExtractedTable(
                source_ref=f'table:{t_idx + 1}',
                headers=headers,
                rows=rows,
                section_heading=current_heading,
            ))
    except Exception as exc:
        logger.error('DOCX content extraction error: %s', exc)
        return ExtractedDocument(
            extractor_type='docx',
            raw_text='\n'.join(paragraphs),
            tables=tables,
            error=f'Partial extraction error: {exc}'
        )

    return ExtractedDocument(
        extractor_type='docx',
        page_count=0,  # DOCX has no fixed page count
        raw_text='\n'.join(paragraphs),
        tables=tables,
        metadata={'paragraph_count': len(paragraphs), 'table_count': len(tables)},
    )
