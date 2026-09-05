"""
XLSX extractor using openpyxl.
"""
import io
import logging
from .base import ExtractedDocument, ExtractedTable

logger = logging.getLogger(__name__)


def _cell_value(cell) -> str:
    """Convert cell value to string safely."""
    if cell.value is None:
        return ''
    return str(cell.value).strip()


def extract_xlsx(file_bytes: bytes) -> ExtractedDocument:
    try:
        import openpyxl
    except ImportError:
        return ExtractedDocument(extractor_type='xlsx', error='openpyxl not installed')

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        return ExtractedDocument(extractor_type='xlsx', error=f'Cannot open XLSX: {exc}')

    tables = []
    sheet_names = []

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_names.append(sheet_name)
            all_rows = []
            for row in ws.iter_rows():
                row_vals = [_cell_value(c) for c in row]
                # Skip completely empty rows
                if any(v for v in row_vals):
                    all_rows.append(row_vals)

            if not all_rows:
                continue

            # First non-empty row is treated as headers
            headers = all_rows[0]
            rows = all_rows[1:]

            tables.append(ExtractedTable(
                source_ref=f'sheet:{sheet_name}',
                headers=headers,
                rows=rows,
                sheet_name=sheet_name,
            ))
    except Exception as exc:
        logger.error('XLSX sheet extraction error: %s', exc)
        return ExtractedDocument(
            extractor_type='xlsx',
            tables=tables,
            error=f'Partial XLSX extraction: {exc}'
        )
    finally:
        wb.close()

    return ExtractedDocument(
        extractor_type='xlsx',
        tables=tables,
        metadata={'sheet_names': sheet_names},
    )
