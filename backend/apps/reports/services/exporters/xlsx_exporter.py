"""
apps.reports.services.exporters.xlsx_exporter — Excel (XLSX) Report Exporter

Uses openpyxl to generate multi-tab, professionally styled Excel workbooks.
"""

import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.reports.models import Report


def generate_xlsx_report(report: Report) -> bytes:
    """
    Generate official XLSX workbook bytes for a given Report instance.
    """
    wb = openpyxl.Workbook()
    content = report.content_json or {}

    navy_fill = PatternFill(start_color="071A3D", end_color="071A3D", fill_type="solid")
    royal_fill = PatternFill(start_color="0B4DB8", end_color="0B4DB8", fill_type="solid")
    light_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    kpi_fill = PatternFill(start_color="EEF2F8", end_color="EEF2F8", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # -------------------------------------------------------------
    # Tab 1: Summary & KPIs
    # -------------------------------------------------------------
    ws_sum = wb.active
    ws_sum.title = "Summary & KPIs"
    ws_sum.views.sheetView[0].showGridLines = True

    # Title & Metadata
    rep_title = content.get('title') or report.title
    ws_sum.merge_cells("A1:F1")
    t_cell = ws_sum["A1"]
    t_cell.value = rep_title
    t_cell.font = Font(name="Calibri", size=16, bold=True, color="071A3D")
    ws_sum.row_dimensions[1].height = 28

    ws_sum["A2"] = "CMPDI · Coal India Limited — AI Mining Intelligence Command Center"
    ws_sum["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    ws_sum["A4"] = "Report Type:"
    ws_sum["B4"] = report.report_type
    ws_sum["C4"] = "Organization:"
    ws_sum["D4"] = report.organization
    ws_sum["E4"] = "Date / Period:"
    ws_sum["F4"] = report.date_range

    ws_sum["A5"] = "Governance Status:"
    ws_sum["B5"] = report.status
    ws_sum["C5"] = "Revision:"
    ws_sum["D5"] = f"#{report.revision_count}"
    ws_sum["E5"] = "Generated At:"
    ws_sum["F5"] = content.get('generated_at') or report.created_at.strftime('%Y-%m-%d %H:%M')

    for r in range(4, 6):
        for c in range(1, 7):
            cell = ws_sum.cell(row=r, column=c)
            cell.font = Font(name="Calibri", size=9, bold=(c % 2 != 0))
            cell.border = thin_border
            if c % 2 != 0:
                cell.fill = light_fill

    # Executive Summary
    ws_sum["A7"] = "EXECUTIVE SUMMARY"
    ws_sum["A7"].font = Font(name="Calibri", size=11, bold=True, color="0B4DB8")
    ws_sum.merge_cells("A8:F10")
    s_cell = ws_sum["A8"]
    s_cell.value = content.get('executive_summary', 'No summary available.')
    s_cell.font = Font(name="Calibri", size=9.5, italic=True)
    s_cell.alignment = Alignment(wrap_text=True, vertical="top")
    s_cell.fill = light_fill
    s_cell.border = thin_border

    # KPIs Grid
    ws_sum["A12"] = "KEY PERFORMANCE INDICATORS"
    ws_sum["A12"].font = Font(name="Calibri", size=11, bold=True, color="0B4DB8")

    kpis = content.get('kpis', [])
    col_idx = 1
    for k in kpis[:4]:
        c_letter = get_column_letter(col_idx)
        val_cell = ws_sum[f"{c_letter}13"]
        lbl_cell = ws_sum[f"{c_letter}14"]
        src_cell = ws_sum[f"{c_letter}15"]

        val_cell.value = k.get('value', '—')
        val_cell.font = Font(name="Calibri", size=14, bold=True, color="0B4DB8")
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell.fill = kpi_fill
        val_cell.border = thin_border

        lbl_cell.value = k.get('label', '')
        lbl_cell.font = Font(name="Calibri", size=9, bold=True, color="1E293B")
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        lbl_cell.fill = kpi_fill
        lbl_cell.border = thin_border

        src_cell.value = f"({k.get('change', 'Verified')})"
        src_cell.font = Font(name="Calibri", size=8, italic=True, color="64748B")
        src_cell.alignment = Alignment(horizontal="center", vertical="center")
        src_cell.fill = kpi_fill
        src_cell.border = thin_border

        col_idx += 1

    # Governance / Sign-off Block
    ws_sum["A17"] = "GOVERNANCE & APPROVAL SIGN-OFF"
    ws_sum["A17"].font = Font(name="Calibri", size=11, bold=True, color="0B4DB8")
    ws_sum["A18"] = "Verification Status"
    ws_sum["B18"] = report.status
    ws_sum["C18"] = "Verified By"
    ws_sum["D18"] = report.verified_by.username if report.verified_by else "Pending"
    ws_sum["E18"] = "Approved By"
    ws_sum["F18"] = report.approved_by.username if report.approved_by else "Pending"

    ws_sum["A19"] = "Revision Number"
    ws_sum["B19"] = str(report.revision_count)
    ws_sum["C19"] = "Verified Date"
    ws_sum["D19"] = report.verified_at.strftime('%Y-%m-%d %H:%M') if report.verified_at else "—"
    ws_sum["E19"] = "Approved Date"
    ws_sum["F19"] = report.approved_at.strftime('%Y-%m-%d %H:%M') if report.approved_at else "—"

    if report.approval_notes:
        ws_sum["A20"] = "Approval Remarks"
        ws_sum["B20"] = report.approval_notes
        ws_sum.merge_cells("B20:F20")

    for r in range(18, 20 + (1 if report.approval_notes else 0)):
        for c in range(1, 7):
            cell = ws_sum.cell(row=r, column=c)
            cell.font = Font(name="Calibri", size=9, bold=(c % 2 != 0))
            cell.border = thin_border
            if c % 2 != 0:
                cell.fill = light_fill

    # -------------------------------------------------------------
    # Tab 2: Report Tables
    # -------------------------------------------------------------
    ws_tbl = wb.create_sheet(title="Report Data")
    ws_tbl.views.sheetView[0].showGridLines = True

    cur_row = 1
    tables = content.get('tables', [])
    for t in tables:
        ws_tbl.cell(row=cur_row, column=1, value=t.get('title', 'Data Table')).font = Font(name="Calibri", size=12, bold=True, color="071A3D")
        cur_row += 1

        cols = t.get('columns', [])
        rows = t.get('rows', [])
        if cols:
            # Table Header
            for c_idx, c_name in enumerate(cols, start=1):
                cell = ws_tbl.cell(row=cur_row, column=c_idx, value=c_name)
                cell.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
                cell.fill = navy_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            ws_tbl.row_dimensions[cur_row].height = 22
            cur_row += 1

            for r_idx, r_vals in enumerate(rows[:50]):
                is_even = r_idx % 2 == 0
                for c_idx, val in enumerate(r_vals, start=1):
                    cell = ws_tbl.cell(row=cur_row, column=c_idx, value=val)
                    cell.font = Font(name="Calibri", size=9)
                    cell.border = thin_border
                    if is_even:
                        cell.fill = light_fill
                cur_row += 1

            if t.get('source'):
                ws_tbl.cell(row=cur_row, column=1, value=f"Source: {t.get('source')}").font = Font(name="Calibri", size=8, italic=True, color="64748B")
                cur_row += 1

            cur_row += 2

    # -------------------------------------------------------------
    # Tab 3: Provenance & Audit
    # -------------------------------------------------------------
    ws_prov = wb.create_sheet(title="Provenance & Audit")
    ws_prov.views.sheetView[0].showGridLines = True

    ws_prov["A1"] = "DATA SOURCE TRACEABILITY & EXTRACTION PROVENANCE"
    ws_prov["A1"].font = Font(name="Calibri", size=12, bold=True, color="071A3D")

    prov_headers = ["Source Type", "Document Title", "Dataset Name", "Sheet Name", "Page No", "Row Index", "Field Name", "Extracted Value", "Method", "Confidence"]
    for c_idx, h in enumerate(prov_headers, start=1):
        cell = ws_prov.cell(row=3, column=c_idx, value=h)
        cell.font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
        cell.fill = royal_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_prov.row_dimensions[3].height = 20

    prov_list = report.provenance_json or []
    for r_idx, p in enumerate(prov_list[:100], start=4):
        ws_prov.cell(row=r_idx, column=1, value=p.get('source_type', 'dataset').upper())
        ws_prov.cell(row=r_idx, column=2, value=p.get('document_title', '—'))
        ws_prov.cell(row=r_idx, column=3, value=p.get('name', '—'))
        ws_prov.cell(row=r_idx, column=4, value=p.get('sheet_name', '—'))
        ws_prov.cell(row=r_idx, column=5, value=p.get('page_number', '—'))
        ws_prov.cell(row=r_idx, column=6, value=p.get('row_index', '—'))
        ws_prov.cell(row=r_idx, column=7, value=p.get('field', '—'))
        ws_prov.cell(row=r_idx, column=8, value=str(p.get('value', '—'))[:60])
        ws_prov.cell(row=r_idx, column=9, value=p.get('extractor_type', '—'))
        ws_prov.cell(row=r_idx, column=10, value=f"{int((p.get('confidence') or 1.0)*100)}%")

        for c_idx in range(1, 11):
            ws_prov.cell(row=r_idx, column=c_idx).border = thin_border
            ws_prov.cell(row=r_idx, column=c_idx).font = Font(name="Calibri", size=8.5)

    # Auto-adjust column widths across all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if '\n' in val:
                    val = max(val.split('\n'), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
