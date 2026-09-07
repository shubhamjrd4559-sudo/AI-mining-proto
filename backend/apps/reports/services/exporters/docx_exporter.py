"""
apps.reports.services.exporters.docx_exporter — Word (DOCX) Report Exporter

Uses python-docx to generate professionally styled Microsoft Word reports.
"""

import io
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from apps.reports.models import Report


def _set_cell_background(cell, hex_color: str):
    """Set background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in dxa."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def generate_docx_report(report: Report) -> bytes:
    """
    Generate official DOCX bytes for a given Report instance.
    """
    doc = DocxDocument()

    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    content = report.content_json or {}
    rep_title = content.get('title') or report.title

    # Header title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(rep_title)
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(7, 26, 61) # Navy
    title_p.paragraph_format.space_after = Pt(2)

    # Subtitle / Agency
    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("CMPDI · Coal India Limited — AI Mining Intelligence Command Center")
    sub_run.font.name = 'Calibri'
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(11, 77, 184) # Royal blue
    sub_p.paragraph_format.space_after = Pt(8)

    # Metadata Bar
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    meta_items = [
        f"Report Type: {report.report_type}",
        f"Organization: {report.organization}",
        f"Period: {report.date_range}",
        f"Status: {report.status}",
        f"Date: {content.get('generated_at') or report.created_at.strftime('%Y-%m-%d')}",
    ]
    meta_run = meta_p.add_run("  |  ".join(meta_items))
    meta_run.font.size = Pt(8.5)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    # Executive Summary Box
    exec_summary = content.get('executive_summary')
    if exec_summary:
        h1 = doc.add_heading("Executive Summary", level=1)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(4)

        box_tbl = doc.add_table(rows=1, cols=1)
        box_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        box_tbl.autofit = False
        cell = box_tbl.cell(0, 0)
        cell.width = Inches(6.8)
        _set_cell_background(cell, "F8FAFC")
        _set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        cp = cell.paragraphs[0]
        crun = cp.add_run(exec_summary)
        crun.font.name = 'Calibri'
        crun.font.size = Pt(10)
        crun.font.italic = True
        crun.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Key Performance Indicators (KPIs)
    kpis = content.get('kpis', [])
    if kpis:
        h_kpi = doc.add_heading("Key Performance Indicators (KPIs)", level=1)
        h_kpi.paragraph_format.space_before = Pt(12)
        h_kpi.paragraph_format.space_after = Pt(6)

        kpi_tbl = doc.add_table(rows=2, cols=len(kpis[:4]))
        kpi_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, k in enumerate(kpis[:4]):
            val_cell = kpi_tbl.cell(0, idx)
            lbl_cell = kpi_tbl.cell(1, idx)
            _set_cell_background(val_cell, "EEF2F8")
            _set_cell_background(lbl_cell, "EEF2F8")
            _set_cell_margins(val_cell, top=80, bottom=40, left=100, right=100)
            _set_cell_margins(lbl_cell, top=20, bottom=80, left=100, right=100)

            vp = val_cell.paragraphs[0]
            vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            vrun = vp.add_run(k.get('value', '—'))
            vrun.font.name = 'Calibri'
            vrun.font.size = Pt(14)
            vrun.font.bold = True
            vrun.font.color.rgb = RGBColor(11, 77, 184)

            lp = lbl_cell.paragraphs[0]
            lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lrun = lp.add_run(f"{k.get('label', '')}\n({k.get('change', '')})")
            lrun.font.name = 'Calibri'
            lrun.font.size = Pt(8)
            lrun.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Tables
    tables = content.get('tables', [])
    for t in tables:
        ht = doc.add_heading(t.get('title', 'Data Table'), level=1)
        ht.paragraph_format.space_before = Pt(14)
        ht.paragraph_format.space_after = Pt(4)

        if t.get('description'):
            dp = doc.add_paragraph(t.get('description'))
            dp.runs[0].font.size = Pt(9)
            dp.runs[0].font.italic = True
            dp.runs[0].font.color.rgb = RGBColor(100, 116, 139)
            dp.paragraph_format.space_after = Pt(6)

        cols = t.get('columns', [])
        rows = t.get('rows', [])
        if cols and rows:
            data_tbl = doc.add_table(rows=len(rows[:25]) + 1, cols=len(cols))
            data_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            # Header
            for c_idx, col_name in enumerate(cols):
                c_cell = data_tbl.cell(0, c_idx)
                _set_cell_background(c_cell, "071A3D")
                _set_cell_margins(c_cell, top=80, bottom=80, left=100, right=100)
                hp = c_cell.paragraphs[0]
                hrun = hp.add_run(col_name)
                hrun.font.name = 'Calibri'
                hrun.font.size = Pt(9)
                hrun.font.bold = True
                hrun.font.color.rgb = RGBColor(255, 255, 255)

            # Data rows
            for r_idx, r_vals in enumerate(rows[:25], start=1):
                bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
                for c_idx, val in enumerate(r_vals):
                    d_cell = data_tbl.cell(r_idx, c_idx)
                    _set_cell_background(d_cell, bg)
                    _set_cell_margins(d_cell, top=60, bottom=60, left=80, right=80)
                    dp = d_cell.paragraphs[0]
                    drun = dp.add_run(str(val))
                    drun.font.name = 'Calibri'
                    drun.font.size = Pt(8.5)
                    drun.font.color.rgb = RGBColor(30, 41, 59)

            if t.get('source'):
                sp = doc.add_paragraph(f"Source: {t.get('source')}")
                sp.runs[0].font.size = Pt(7.5)
                sp.runs[0].font.italic = True
                sp.runs[0].font.color.rgb = RGBColor(100, 116, 139)
                sp.paragraph_format.space_after = Pt(8)

    # Operational Analysis
    analysis = content.get('analysis')
    if analysis:
        ha = doc.add_heading("Operational & Geological Analysis", level=1)
        ha.paragraph_format.space_before = Pt(14)
        ha.paragraph_format.space_after = Pt(4)
        for p in analysis.split('\n\n'):
            if p.strip():
                para = doc.add_paragraph(p.strip())
                para.paragraph_format.space_after = Pt(6)
                para.runs[0].font.size = Pt(9.5)

    # Conclusions
    conclusions = content.get('conclusions')
    if conclusions:
        hc = doc.add_heading("Strategic Conclusions & Directives", level=1)
        hc.paragraph_format.space_before = Pt(14)
        hc.paragraph_format.space_after = Pt(4)
        for line in conclusions.split('\n'):
            if line.strip():
                p = doc.add_paragraph(line.strip())
                p.paragraph_format.space_after = Pt(4)
                p.runs[0].font.size = Pt(9.5)

    # Source Provenance Table
    provenance = report.provenance_json or []
    if provenance:
        hp = doc.add_heading("Source Traceability & Audit Provenance", level=1)
        hp.paragraph_format.space_before = Pt(14)
        hp.paragraph_format.space_after = Pt(4)

        p_tbl = doc.add_table(rows=len(provenance[:15]) + 1, cols=5)
        p_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Source Type", "Document / Dataset", "Sheet / Page", "Row / Ref", "Method"]
        for c_idx, h in enumerate(headers):
            cell = p_tbl.cell(0, c_idx)
            _set_cell_background(cell, "334155")
            _set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            hp = cell.paragraphs[0]
            hrun = hp.add_run(h)
            hrun.font.bold = True
            hrun.font.size = Pt(8.5)
            hrun.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, p in enumerate(provenance[:15], start=1):
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            p_type = p.get('source_type', 'dataset').upper()
            name = p.get('name') or p.get('document_title') or 'Dataset'
            pg_sheet = f"p. {p.get('page_number')}" if p.get('page_number') else (p.get('sheet_name') or 'Main')
            row_ref = p.get('source_ref') or f"Row {p.get('row_index', '—')}"
            method = f"{p.get('extractor_type', 'doc')} ({int((p.get('confidence') or 1.0)*100)}%)"

            vals = [p_type, name[:30], str(pg_sheet), str(row_ref), method]
            for c_idx, val in enumerate(vals):
                cell = p_tbl.cell(r_idx, c_idx)
                _set_cell_background(cell, bg)
                _set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
                cp = cell.paragraphs[0]
                crun = cp.add_run(val)
                crun.font.size = Pt(8)
                crun.font.color.rgb = RGBColor(30, 41, 59)

    # Sign-off Block
    doc.add_heading("Governance & Authority Sign-off", level=1)
    s_tbl = doc.add_table(rows=2, cols=3)
    s_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_data = [
        [f"Verification Status: {report.status}", f"Verified By: {report.verified_by.username if report.verified_by else 'Pending'}", f"Approved By: {report.approved_by.username if report.approved_by else 'Pending'}"],
        [f"Revision: #{report.revision_count}", f"Verified At: {report.verified_at.strftime('%Y-%m-%d %H:%M') if report.verified_at else '—'}", f"Approved At: {report.approved_at.strftime('%Y-%m-%d %H:%M') if report.approved_at else '—'}"]
    ]
    for r_idx, r_list in enumerate(s_data):
        for c_idx, val in enumerate(r_list):
            cell = s_tbl.cell(r_idx, c_idx)
            _set_cell_background(cell, "F1F5F9")
            _set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(8.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
